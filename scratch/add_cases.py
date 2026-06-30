import openpyxl

file_path = r"c:\Users\xkait\Desktop\Taller\Catch-Go\Documentacion\Unidad #3\3.1.3 Planilla Casos de Prueba.xlsx"
wb = openpyxl.load_workbook(file_path)

# 1. Modificar Casos de Prueba
sheet = wb['Casos de Prueba']

new_cases_data = {
    70: [9, 'Navegación y Menús', 'frontend', 'Prueba Unitarias', 23, 'Prueba unitaria de Barra de Navegación (Navbar)', 
         '1. Renderizar Navbar en estado sin sesión.', 'localStorage vacío, ruta /', 
         'Se visualizan enlaces de Quiénes Somos, Contacto, Términos y botones de ingreso.', 'OK'],
    71: [None, None, None, None, None, None, 
         '2. Renderizar Navbar con sesión activa de Trabajador o Empresa.', 'localStorage con user_info de TRABAJADOR o EMPRESA', 
         'Se visualizan opciones de dashboard específicas por rol y campana de notificaciones.', 'OK'],
    72: [None, None, None, None, None, None, 
         '3. Ejecutar flujo de cierre de sesión.', 'Clic en salir y confirmar', 
         'Se limpian claves de localStorage y se redirige a /login.', 'OK'],
         
    74: [9, 'Navegación y Menús', 'frontend', 'Prueba Unitarias', 24, 'Prueba unitaria de Campana de Notificaciones', 
         '1. Renderizar campana con unreadCount > 0.', 'unreadCount = 1', 
         'Se dibuja un círculo rojo con el número de notificaciones no leídas.', 'OK'],
    75: [None, None, None, None, None, None, 
         '2. Hacer clic en campana y desplegar listado.', 'Clic en botón campana', 
         'Se abre el panel mostrando las notificaciones y botones de Limpiar/Marcar como leída.', 'OK'],
    76: [None, None, None, None, None, None, 
         '3. Ejecutar limpieza o marcado de lectura.', 'Clic en Limpiar o en notificación', 
         'Se invocan las funciones del contexto clearAll o markAsRead correspondientes.', 'OK'],
         
    78: [10, 'Pagos', 'frontend', 'Prueba Unitarias', 25, 'Prueba unitaria de Modal de Pago (PaymentGatewayModal)', 
         '1. Renderizar modal en transferencia manual sin archivo adjunto.', 'isOpen = true, file = null', 
         'El botón Confirmar Pago está deshabilitado.', 'OK'],
    79: [None, None, None, None, None, None, 
         '2. Adjuntar comprobante de pago en transferencia manual.', 'Archivo comprobante.pdf', 
         'El botón Confirmar Pago se habilita y al pulsar llama a onSubmit con los datos.', 'OK'],
    80: [None, None, None, None, None, None, 
         '3. Cambiar a pasarela Webpay Plus y presionar pagar.', 'Clic en pestaña Webpay y botón pagar', 
         'Se invoca la función onPayWithWebpay y el botón muestra estado de carga.', 'OK'],
         
    82: [2, 'Iniciar sesión', 'frontend', 'Prueba Unitarias', 26, 'Prueba unitaria de Página de Login (page.tsx)', 
         '1. Intentar iniciar sesión con campos vacíos.', 'email = "", password = ""', 
         'Se despliegan mensajes locales indicando que los campos son obligatorios.', 'OK'],
    83: [None, None, None, None, None, None, 
         '2. Iniciar sesión exitosamente mediante API.', 'Credenciales correctas de Trabajador / Empresa', 
         'Guarda auth_token e user_info en localStorage y redirige al perfil correspondiente.', 'OK'],
    84: [None, None, None, None, None, None, 
         '3. Iniciar sesión con error de credenciales.', 'Credenciales erróneas', 
         'Se despliega el mensaje de error retornado por la API de autenticación.', 'OK']
}

for row_num, row_data in new_cases_data.items():
    for col_idx, val in enumerate(row_data, start=1):
        sheet.cell(row=row_num, column=col_idx, value=val)
        
print("Casos de Prueba actualizados en Excel.")

# 2. Modificar Totales
totals_sheet = wb['Totales']
# Celda C11 (Fila 11 Columna 3) es Total casos de Prueba
totals_sheet.cell(row=11, column=3, value=26)
# Celda C12 (Fila 12 Columna 3) es Total casos de Prueba OK
totals_sheet.cell(row=12, column=3, value=26)

print("Totales actualizados en Excel.")

wb.save(file_path)
print("Archivo guardado con éxito.")
