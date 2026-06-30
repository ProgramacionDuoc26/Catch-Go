import docx

doc_path = r"c:\Users\xkait\Desktop\Taller\Catch-Go\Documentacion\Unidad #3\3.1.2 Plan Pruebas Funcionales.docx"
doc = docx.Document(doc_path)

# --- 1. ACTUALIZAR TABLA 3 (TRAZABILIDAD DE CASOS DE PRUEBAS - REQUISITOS) ---
t3 = doc.tables[3]
new_t3_rows = [
    ['CP-23', '', '', '', '', '', '', 'X', ''],
    ['CP-24', '', '', '', '', '', '', '', 'X'],
    ['CP-25', '', '', '', '', 'X', '', '', ''],
    ['CP-26', '', 'X', '', '', '', '', '', '']
]
for row_data in new_t3_rows:
    row = t3.add_row()
    for col_idx, val in enumerate(row_data):
        row.cells[col_idx].text = val

print("Tabla 3 (Trazabilidad) actualizada en Word.")

# --- 2. ACTUALIZAR TABLA 5 (ESTRATEGIA DE EJECUCIÓN DE PRUEBAS) ---
t5 = doc.tables[5]
new_t5_rows = [
    ['CP-23', 'X', '', ''],
    ['CP-24', 'X', '', ''],
    ['CP-25', 'X', '', ''],
    ['CP-26', 'X', '', '']
]
for row_data in new_t5_rows:
    row = t5.add_row()
    for col_idx, val in enumerate(row_data):
        row.cells[col_idx].text = val

print("Tabla 5 (Estrategia) actualizada en Word.")

# --- 3. ACTUALIZAR TABLA 4 (DEFINICIÓN DE LOS CASOS DE PRUEBAS) ---
t4 = doc.tables[4]

# Estructura de cada caso: 7 filas
# Fila 1: [Título del caso, ID, ID]
# Fila 2: [Título del caso, "¿Prueba de despliegue?", "No"]
# Fila 3: [Descripción, Descripción, Descripción]
# Fila 4: [Prerrequisitos, Prerrequisitos, Prerrequisitos]
# Fila 5: [Pasos, Pasos, Pasos]
# Fila 6: [Resultado esperado, Resultado esperado, Resultado esperado]
# Fila 7: [Resultado obtenido, Resultado obtenido, Resultado obtenido]

new_t4_definitions = [
    # CP-23
    [
        ['Prueba unitaria de Barra de Navegación (Navbar)', 'CP-23', 'CP-23'],
        ['Prueba unitaria de Barra de Navegación (Navbar)', '¿Prueba de despliegue?', 'No'],
        ['Descripción: Verificar el renderizado de opciones según el rol del usuario (Público, Trabajador, Empresa, Administrador) y flujo de cierre de sesión.', 
         'Descripción: Verificar el renderizado de opciones según el rol del usuario (Público, Trabajador, Empresa, Administrador) y flujo de cierre de sesión.', 
         'Descripción: Verificar el renderizado de opciones según el rol del usuario (Público, Trabajador, Empresa, Administrador) y flujo de cierre de sesión.'],
        ['Prerrequisitos: Componente renderizado bajo React Testing Library con mocks de enrutamiento y localStorage.', 
         'Prerrequisitos: Componente renderizado bajo React Testing Library con mocks de enrutamiento y localStorage.', 
         'Prerrequisitos: Componente renderizado bajo React Testing Library con mocks de enrutamiento y localStorage.'],
        ['Pasos: 1. Renderizar Navbar sin sesión activa. 2. Cambiar a rol Trabajador/Empresa/Admin y verificar links de panel. 3. Hacer clic en "Cerrar Sesión" y confirmar.', 
         'Pasos: 1. Renderizar Navbar sin sesión activa. 2. Cambiar a rol Trabajador/Empresa/Admin y verificar links de panel. 3. Hacer clic en "Cerrar Sesión" y confirmar.', 
         'Pasos: 1. Renderizar Navbar sin sesión activa. 2. Cambiar a rol Trabajador/Empresa/Admin y verificar links de panel. 3. Hacer clic en "Cerrar Sesión" y confirmar.'],
        ['Resultado esperado: Se muestran links públicos sin sesión, menús de dashboard correspondientes para Trabajador/Empresa/Admin, y se limpia la sesión en el logout.', 
         'Resultado esperado: Se muestran links públicos sin sesión, menús de dashboard correspondientes para Trabajador/Empresa/Admin, y se limpia la sesión en el logout.', 
         'Resultado esperado: Se muestran links públicos sin sesión, menús de dashboard correspondientes para Trabajador/Empresa/Admin, y se limpia la sesión en el logout.'],
        ['Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.']
    ],
    # CP-24
    [
        ['Prueba unitaria de Campana de Notificaciones', 'CP-24', 'CP-24'],
        ['Prueba unitaria de Campana de Notificaciones', '¿Prueba de despliegue?', 'No'],
        ['Descripción: Validación del despliegue de alertas pendientes, contador de mensajes no leídos, y acciones de marcar como leída y limpiar.', 
         'Descripción: Validación del despliegue de alertas pendientes, contador de mensajes no leídos, y acciones de marcar como leída y limpiar.', 
         'Descripción: Validación del despliegue de alertas pendientes, contador de mensajes no leídos, y acciones de marcar como leída y limpiar.'],
        ['Prerrequisitos: Contexto de NotificationContext simulado.', 
         'Prerrequisitos: Contexto de NotificationContext simulado.', 
         'Prerrequisitos: Contexto de NotificationContext simulado.'],
        ['Pasos: 1. Renderizar campana con 1 notificación activa. 2. Hacer clic en campana para abrir panel. 3. Hacer clic en "Limpiar" o "Marcar como leída".', 
         'Pasos: 1. Renderizar campana con 1 notificación activa. 2. Hacer clic en campana para abrir panel. 3. Hacer clic en "Limpiar" o "Marcar como leída".', 
         'Pasos: 1. Renderizar campana con 1 notificación activa. 2. Hacer clic en campana para abrir panel. 3. Hacer clic en "Limpiar" o "Marcar como leída".'],
        ['Resultado esperado: Muestra la burbuja con el número de notificaciones no leídas, abre el panel de lista, e invoca clearAll / markAsRead al interactuar.', 
         'Resultado esperado: Muestra la burbuja con el número de notificaciones no leídas, abre el panel de lista, e invoca clearAll / markAsRead al interactuar.', 
         'Resultado esperado: Muestra la burbuja con el número de notificaciones no leídas, abre el panel de lista, e invoca clearAll / markAsRead al interactuar.'],
        ['Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.']
    ],
    # CP-25
    [
        ['Prueba unitaria de Modal de Pago (PaymentGatewayModal)', 'CP-25', 'CP-25'],
        ['Prueba unitaria de Modal de Pago (PaymentGatewayModal)', '¿Prueba de despliegue?', 'No'],
        ['Descripción: Validación de la obligatoriedad de comprobante de pago en transferencia manual y flujo alternativo con pasarela Webpay Plus.', 
         'Descripción: Validación de la obligatoriedad de comprobante de pago en transferencia manual y flujo alternativo con pasarela Webpay Plus.', 
         'Descripción: Validación de la obligatoriedad de comprobante de pago en transferencia manual y flujo alternativo con pasarela Webpay Plus.'],
        ['Prerrequisitos: Mocks de FileReader y temporizadores (fake timers).', 
         'Prerrequisitos: Mocks de FileReader y temporizadores (fake timers).', 
         'Prerrequisitos: Mocks de FileReader y temporizadores (fake timers).'],
        ['Pasos: 1. Abrir modal y presionar confirmar sin archivo. 2. Seleccionar archivo y presionar confirmar. 3. Ir a pestaña Webpay y presionar pagar.', 
         'Pasos: 1. Abrir modal y presionar confirmar sin archivo. 2. Seleccionar archivo y presionar confirmar. 3. Ir a pestaña Webpay y presionar pagar.', 
         'Pasos: 1. Abrir modal y presionar confirmar sin archivo. 2. Seleccionar archivo y presionar confirmar. 3. Ir a pestaña Webpay y presionar pagar.'],
        ['Resultado esperado: El botón de confirmar está deshabilitado si no hay archivo, al subir el archivo se habilita e invoca onSubmit, y al usar Webpay llama a onPayWithWebpay.', 
         'Resultado esperado: El botón de confirmar está deshabilitado si no hay archivo, al subir el archivo se habilita e invoca onSubmit, y al usar Webpay llama a onPayWithWebpay.', 
         'Resultado esperado: El botón de confirmar está deshabilitado si no hay archivo, al subir el archivo se habilita e invoca onSubmit, y al usar Webpay llama a onPayWithWebpay.'],
        ['Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.']
    ],
    # CP-26
    [
        ['Prueba unitaria de Página de Login (page.tsx)', 'CP-26', 'CP-26'],
        ['Prueba unitaria de Página de Login (page.tsx)', '¿Prueba de despliegue?', 'No'],
        ['Descripción: Verificación de validación de campos obligatorios locales, redireccionamientos correctos por rol en login exitoso, y renderizado de errores.', 
         'Descripción: Verificación de validación de campos obligatorios locales, redireccionamientos correctos por rol en login exitoso, y renderizado de errores.', 
         'Descripción: Verificación de validación de campos obligatorios locales, redireccionamientos correctos por rol en login exitoso, y renderizado de errores.'],
        ['Prerrequisitos: Mocks de authApi.login, useRouter y supabase.auth.signOut.', 
         'Prerrequisitos: Mocks de authApi.login, useRouter y supabase.auth.signOut.', 
         'Prerrequisitos: Mocks de authApi.login, useRouter y supabase.auth.signOut.'],
        ['Pasos: 1. Presionar "Ingresar ahora" con campos vacíos. 2. Llenar credenciales de trabajador y presionar enviar. 3. Simular credenciales erróneas.', 
         'Pasos: 1. Presionar "Ingresar ahora" con campos vacíos. 2. Llenar credenciales de trabajador y presionar enviar. 3. Simular credenciales erróneas.', 
         'Pasos: 1. Presionar "Ingresar ahora" con campos vacíos. 2. Llenar credenciales de trabajador y presionar enviar. 3. Simular credenciales erróneas.'],
        ['Resultado esperado: Muestra mensajes de campos obligatorios, guarda auth_token en localStorage y redirige al dashboard por rol al tener éxito, y muestra el error global si falla.', 
         'Resultado esperado: Muestra mensajes de campos obligatorios, guarda auth_token en localStorage y redirige al dashboard por rol al tener éxito, y muestra el error global si falla.', 
         'Resultado esperado: Muestra mensajes de campos obligatorios, guarda auth_token en localStorage y redirige al dashboard por rol al tener éxito, y muestra el error global si falla.'],
        ['Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.', 
         'Resultado obtenido: Ejecutado con éxito. Pruebas automatizadas en Vitest pasadas al 100% sin advertencias.']
    ]
]

for case_definition in new_t4_definitions:
    for row_data in case_definition:
        row = t4.add_row()
        for col_idx, val in enumerate(row_data):
            row.cells[col_idx].text = val

print("Tabla 4 (Definición de Casos de Pruebas) actualizada en Word.")

doc.save(doc_path)
print("Documento de Word guardado con éxito.")
