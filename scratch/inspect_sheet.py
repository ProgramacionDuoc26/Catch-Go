import openpyxl

wb = openpyxl.load_workbook(r"c:\Users\xkait\Desktop\Taller\Catch-Go\Documentacion\Unidad #3\3.1.3 Planilla Casos de Prueba.xlsx")
sheet = wb['Casos de Prueba']

# Imprimir las columnas (Fila 12)
header = [cell for cell in next(sheet.iter_rows(min_row=12, max_row=12, values_only=True))]
print("Headers:", header)

# Leer todos los casos
cases = []
for r_idx, row in enumerate(sheet.iter_rows(min_row=13, values_only=True), start=13):
    cp_num = row[4] # Columna E (N° CASO de PRUEBA)
    if cp_num is not None:
        cases.append((r_idx, cp_num, row[5], row[2])) # Row Index, CP Num, Title, Module

print(f"Total cases found: {len(cases)}")
for case in cases:
    print(f"Row {case[0]}: CP-{case[1]} - {case[2]} ({case[3]})")
